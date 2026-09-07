import csv
import hashlib
import io
import re
from dataclasses import dataclass
from pathlib import PurePath

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader

from backend.rag.loader import SourceDocument


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown", ".csv", ".xlsx"}
MAX_UPLOAD_BYTES = 10 * 1024 * 1024


class DocumentLoadError(ValueError):
    """A controlled error raised when an uploaded document cannot be loaded."""


@dataclass(frozen=True)
class UploadedDocument:
    document: SourceDocument
    filename: str
    extension: str
    content_type: str | None
    section_count: int
    byte_count: int


def _safe_filename(filename: str) -> str:
    name = PurePath(filename or "").name
    if not name or name in {".", ".."}:
        raise DocumentLoadError("a filename is required")
    return name


def _normalise_text(value: str) -> str:
    return re.sub(r"\n{3,}", "\n\n", re.sub(r"[ \t]+", " ", value)).strip()


def _decode_text(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentLoadError("text files must be valid UTF-8") from exc


def _tabular_text(rows: list[list[object]], label: str = "Row") -> tuple[str, int]:
    if not rows:
        return "", 0
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    headers = [value or f"Column {index + 1}" for index, value in enumerate(headers)]
    lines = []
    for row_number, row in enumerate(rows[1:], start=2):
        values = list(row) + [None] * max(0, len(headers) - len(row))
        fields = [
            f"{header}: {str(value).strip()}"
            for header, value in zip(headers, values)
            if value is not None and str(value).strip()
        ]
        if fields:
            lines.append(f"{label} {row_number}: " + " | ".join(fields))
    return "\n".join(lines), len(lines)


class UploadedDocumentLoader:
    def load(
        self,
        data: bytes,
        filename: str,
        content_type: str | None = None,
    ) -> UploadedDocument:
        if not data:
            raise DocumentLoadError("uploaded file is empty")
        if len(data) > MAX_UPLOAD_BYTES:
            raise DocumentLoadError("uploaded file exceeds the 10 MB size limit")

        safe_name = _safe_filename(filename)
        extension = PurePath(safe_name).suffix.lower()
        if extension not in SUPPORTED_EXTENSIONS:
            raise DocumentLoadError(
                "unsupported file type; supported formats are PDF, DOCX, TXT, MD, CSV, and XLSX"
            )

        try:
            content, section_count = self._extract(data, extension)
        except DocumentLoadError:
            raise
        except Exception as exc:
            raise DocumentLoadError(f"could not read {safe_name}") from exc

        content = _normalise_text(content)
        if not content:
            raise DocumentLoadError("uploaded file contains no readable text")

        source_id = "upload-" + hashlib.sha256(safe_name.encode() + data).hexdigest()[:16]
        document = SourceDocument(
            source_id=source_id,
            citation_id=None,
            source_title=safe_name,
            source_url=f"uploaded://{source_id}/{safe_name}",
            content=content,
            metadata={
                "source_type": "uploaded_document",
                "filename": safe_name,
                "file_extension": extension,
                "mime_type": content_type,
            },
        )
        return UploadedDocument(
            document=document,
            filename=safe_name,
            extension=extension,
            content_type=content_type,
            section_count=section_count,
            byte_count=len(data),
        )

    def _extract(self, data: bytes, extension: str) -> tuple[str, int]:
        if extension in {".txt", ".md", ".markdown"}:
            text = _decode_text(data)
            return text, max(1, len([line for line in text.splitlines() if line.strip()]))
        if extension == ".pdf":
            reader = PdfReader(io.BytesIO(data))
            pages = [page.extract_text() or "" for page in reader.pages]
            text = "\n\n".join(
                f"Page {index}:\n{page}"
                for index, page in enumerate(pages, start=1)
                if page.strip()
            )
            if not text:
                raise DocumentLoadError("PDF contains no readable text")
            return text, len([page for page in pages if page.strip()])
        if extension == ".docx":
            document = DocxDocument(io.BytesIO(data))
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    values = [cell.text.strip() for cell in row.cells]
                    if any(values):
                        parts.append(" | ".join(values))
            return "\n".join(parts), len(parts)
        if extension == ".csv":
            rows = list(csv.reader(io.StringIO(_decode_text(data))))
            return _tabular_text(rows)
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts = []
        count = 0
        for worksheet in workbook.worksheets:
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            text, row_count = _tabular_text(rows, label=f"{worksheet.title} row")
            if text:
                parts.append(f"Worksheet: {worksheet.title}\n{text}")
                count += row_count
        return "\n\n".join(parts), count